import asyncio
from pathlib import Path
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches

HTML_FILES = [f"{i}.html" for i in range(1, 9)]
IMG_DIR = Path("slide_imgs")
IMG_DIR.mkdir(exist_ok=True)

async def screenshot_slides():
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1280, "height": 720})
        for html in HTML_FILES:
            path = Path(html).resolve()
            await page.goto(f"file://{path}")
            await page.wait_for_timeout(500)
            out = IMG_DIR / (Path(html).stem + ".png")
            await page.screenshot(path=str(out), clip={"x":0,"y":0,"width":1280,"height":720})
            print(f"Captured {html} → {out}")
        await browser.close()

def build_pptx():
    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for i in range(1, 9):
        img = IMG_DIR / f"{i}.png"
        sl = prs.slides.add_slide(blank)
        sl.shapes.add_picture(str(img), 0, 0, Inches(13.33), Inches(7.5))
        print(f"Added slide {i}")
    prs.save("resume_analyzer.pptx")
    print("Done! → resume_analyzer.pptx")

asyncio.run(screenshot_slides())
build_pptx()