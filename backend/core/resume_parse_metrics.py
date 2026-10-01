from io import BytesIO
from pathlib import Path
import re

import pymupdf
from docx import Document

from core.schemas import ResumeParseReport


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
}

MIN_PDF_PAGE_CHARS = 20


def _clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _parse_pdf(
    file_bytes: bytes,
    filename: str,
) -> ResumeParseReport:
    parsed_pages = 0
    characters_extracted = 0

    with pymupdf.open(
        stream=file_bytes,
        filetype="pdf",
    ) as document:
        total_pages = len(document)

        for page in document:
            text = _clean_text(
                page.get_text("text")
            )

            characters_extracted += len(text)

            if len(text) >= MIN_PDF_PAGE_CHARS:
                parsed_pages += 1

    percentage = (
        round(
            parsed_pages
            / total_pages
            * 100
        )
        if total_pages
        else 0
    )

    warning = None

    if percentage < 100:
        warning = (
            "Some PDF pages contained little or no "
            "extractable text. OCR may be required."
        )

    return ResumeParseReport(
        filename=filename,
        file_type="pdf",
        parse_percentage=percentage,
        parsed_units=parsed_pages,
        total_units=total_pages,
        unit="pages",
        characters_extracted=characters_extracted,
        warning=warning,
    )


def _parse_docx(
    file_bytes: bytes,
    filename: str,
) -> ResumeParseReport:
    document = Document(
        BytesIO(file_bytes)
    )

    parts = []

    for paragraph in document.paragraphs:
        parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)

    text = _clean_text(
        "\n".join(parts)
    )

    parsed = bool(text)

    return ResumeParseReport(
        filename=filename,
        file_type="docx",
        parse_percentage=100 if parsed else 0,
        parsed_units=1 if parsed else 0,
        total_units=1,
        unit="document",
        characters_extracted=len(text),
        warning=None if parsed else "No extractable text found.",
    )


def _parse_txt(
    file_bytes: bytes,
    filename: str,
) -> ResumeParseReport:
    try:
        raw_text = file_bytes.decode("utf-8")
    except UnicodeDecodeError as error:
        raise ValueError(
            "TXT file must use UTF-8 encoding."
        ) from error

    text = _clean_text(raw_text)
    parsed = bool(text)

    return ResumeParseReport(
        filename=filename,
        file_type="txt",
        parse_percentage=100 if parsed else 0,
        parsed_units=1 if parsed else 0,
        total_units=1,
        unit="document",
        characters_extracted=len(text),
        warning=None if parsed else "No extractable text found.",
    )


def inspect_resume_parse(
    file_bytes: bytes,
    filename: str,
) -> ResumeParseReport:
    extension = Path(
        filename
    ).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Only PDF, DOCX and TXT files are supported."
        )

    if not file_bytes:
        raise ValueError(
            "Uploaded resume is empty."
        )

    if extension == ".pdf":
        return _parse_pdf(
            file_bytes,
            filename,
        )

    if extension == ".docx":
        return _parse_docx(
            file_bytes,
            filename,
        )

    return _parse_txt(
        file_bytes,
        filename,
    )
